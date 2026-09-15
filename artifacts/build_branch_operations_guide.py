from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "SAT-J_Ent_Branch_Operations_Guide.docx"

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F7FA"
MID_GRAY = "666666"
LINE_GRAY = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), LINE_GRAY)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def remove_paragraph_border(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_header(hdr)
    prevent_row_split(hdr)
    for idx, value in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = widths[idx]
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        if row_idx % 2:
            for cell in cells:
                set_cell_shading(cell, PALE_BLUE)
        for idx, value in enumerate(values):
            cell = cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = widths[idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.add_run(str(value))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.85)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.7)
normal.font.color.rgb = RGBColor(34, 34, 34)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

for style_name, size, before, after in (
    ("Title", 25, 0, 12),
    ("Heading 1", 17, 15, 7),
    ("Heading 2", 13, 10, 5),
    ("Heading 3", 11, 8, 4),
):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

remove_paragraph_border(styles["Title"])

for list_style in ("List Bullet", "List Bullet 2", "List Number"):
    styles[list_style].font.name = "Aptos"
    styles[list_style].font.size = Pt(10.5)
    styles[list_style].paragraph_format.space_after = Pt(3)
    styles[list_style].paragraph_format.line_spacing = 1.08

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("SAT-J Ent Branch Operations Guide   |   ")
footer_run.font.name = "Aptos"
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor.from_string(MID_GRAY)
add_page_number(footer)

title = doc.add_paragraph(style="Title")
title.add_run("SAT J Ent Branch Operations Guide")
remove_paragraph_border(title)
subtitle = doc.add_paragraph()
subtitle_run = subtitle.add_run("How employees use the system at each branch")
subtitle_run.bold = True
subtitle_run.font.size = Pt(13)
subtitle_run.font.color.rgb = RGBColor.from_string(MID_GRAY)
subtitle.paragraph_format.space_after = Pt(14)

intro = doc.add_paragraph()
intro.add_run("Purpose  ").bold = True
intro.add_run(
    "This guide explains how SAT-J Ent will operate from each branch, who performs each task, "
    "how transactions affect branch records, and how head office monitors the company. All branches "
    "use the same web application and central database. The employee's role and active branch assignments "
    "determine what the employee can see and do."
)

add_heading(doc, "Operating Model", 1)
doc.add_paragraph(
    "Every branch works in the same system. Product, supplier, customer, and company reference data are shared, "
    "while operational records such as stock, sales, purchases, payments, and reports retain their branch identity. "
    "A transaction completed at one branch becomes available to authorised management immediately."
)
add_table(
    doc,
    ["Area", "How it works"],
    [
        ["Single application", "Employees use the SAT-J Ent web application on an authorised phone, tablet, or computer."],
        ["Central database", "All branches write to the same company database, which supports current management reporting."],
        ["Role permissions", "The assigned role controls the modules and actions available to each employee."],
        ["Branch assignment", "Branch employees can work only with records belonging to their active assigned branches."],
        ["Company oversight", "Owner and Super Admin accounts can review authorised information across all branches."],
    ],
    [Inches(1.65), Inches(5.0)],
)

add_heading(doc, "Branch Setup", 1)
doc.add_paragraph("Before a branch starts daily operations, a Super Admin completes the following setup:")
for item in (
    "Create the branch and confirm that it is active.",
    "Create an employee profile for each staff member.",
    "Assign the appropriate system role to each employee.",
    "Assign branch-level employees to one or more branches.",
    "Activate employee access and provide individual sign-in credentials.",
    "Set up the product catalogue, selling prices, suppliers, and opening stock required for the branch.",
):
    add_number(doc, item)
doc.add_paragraph(
    "Employees should use their own accounts. Shared logins weaken the audit trail because the system records the signed-in user for important operations."
)

add_heading(doc, "Roles Used at a Branch", 1)
add_table(
    doc,
    ["Role", "Primary responsibility", "Operating scope"],
    [
        ["Branch Manager", "Oversees sales, purchasing, stock, payments, transfers, quotations, and branch reporting.", "Assigned branches"],
        ["Sales or Cashier", "Handles customers, sales, collections, receivables, and quotation follow-up.", "Assigned branches"],
        ["Inventory Officer", "Maintains stock, receives goods, counts stock, records adjustments, and handles transfers.", "Assigned branches"],
        ["Owner or Management", "Reviews company performance, branch comparisons, finances, inventory, and audit history.", "Company-wide"],
        ["Super Admin", "Manages branches and employee access and has full operational oversight.", "Company-wide"],
    ],
    [Inches(1.45), Inches(3.85), Inches(1.35)],
)

add_heading(doc, "Branch Manager Responsibilities", 2)
doc.add_paragraph("The Branch Manager coordinates the complete operation of an assigned branch. The role can:")
for item in (
    "Review the branch dashboard, operational queues, and branch reports.",
    "Create and update suppliers, purchases, and supplier payments within authorised scope.",
    "Receive purchased goods into branch inventory and cancel eligible purchases.",
    "Create and update customers, complete sales, apply discounts, override prices with a reason, and cancel eligible completed sales.",
    "Record customer payments, reverse posted payments when authorised, and monitor receivables.",
    "Post opening stock, adjust stock, complete physical counts, and maintain minimum stock levels.",
    "Create, approve, dispatch, receive, or cancel eligible branch transfers.",
    "Review quotation requests for the branch and update their status.",
    "Export authorised branch reports.",
):
    add_bullet(doc, item)
doc.add_paragraph(
    "The Branch Manager cannot manage system roles, company branches, or employee access unless a separate company-level account grants those permissions."
)

doc.add_page_break()
add_heading(doc, "Sales and Cashier Responsibilities", 2)
doc.add_paragraph("The Sales or Cashier role handles customer-facing transactions. The role can:")
for item in (
    "View the catalogue, prices, and available stock for assigned branches.",
    "Create or update customer records.",
    "Create and complete sales for an assigned branch.",
    "Record customer payments and monitor outstanding receivables.",
    "View stock transfers involving the assigned branch.",
    "Read quotation requests and update their status within branch scope.",
):
    add_bullet(doc, item)
doc.add_paragraph(
    "This role does not perform stock adjustments, receive purchases, approve transfers, cancel completed sales, or override selling prices under the standard permission set."
)

add_heading(doc, "Inventory Officer Responsibilities", 2)
doc.add_paragraph("The Inventory Officer maintains the physical stock position of an assigned branch. The role can:")
for item in (
    "Post opening stock when a product begins inventory tracking at the branch.",
    "View current quantities, minimum levels, and the stock movement history.",
    "Create and complete authorised stock adjustments.",
    "Start, enter, and reconcile physical stock counts.",
    "Maintain branch-specific minimum stock levels.",
    "View suppliers and receive goods against authorised purchases.",
    "Create and dispatch transfers from the branch and receive transfers at the destination branch.",
    "Monitor low-stock, out-of-stock, and in-transit items on the branch dashboard.",
):
    add_bullet(doc, item)
doc.add_paragraph(
    "The Inventory Officer can view sales because completed sales affect stock, but the standard role does not operate the cashier workflow or approve and cancel transfers."
)

add_heading(doc, "Daily Sales Workflow", 1)
doc.add_paragraph("A cashier or authorised manager processes a normal branch sale in this order:")
sales_steps = (
    ("Select the branch", "The sale belongs to one authorised branch. A branch employee cannot submit another branch."),
    ("Select the customer", "Use an existing customer, create a new customer, or use the walk-in customer where appropriate."),
    ("Add products", "Select product variants and enter the quantities being sold."),
    ("Resolve the price", "The system uses the applicable company or branch price. An authorised override requires a reason."),
    ("Complete the sale", "Completion records the sale and reduces the branch inventory through stock movements."),
    ("Record payment", "Capture the amount received and payment method. Additional payments can be recorded later."),
    ("Track the balance", "Any unpaid amount remains on the customer account as a receivable."),
)
add_table(doc, ["Step", "Branch action", "System result"], [[str(i + 1), name, result] for i, (name, result) in enumerate(sales_steps)], [Inches(0.55), Inches(1.55), Inches(4.55)])

add_heading(doc, "Daily Purchasing Workflow", 1)
doc.add_paragraph("A Branch Manager or another authorised employee processes purchases as follows:")
purchase_steps = (
    ("Select the supplier", "Use an existing supplier or create an authorised supplier record."),
    ("Create the purchase", "Choose the branch and add ordered products, quantities, and purchase costs."),
    ("Place the order", "The order can be tracked without changing stock."),
    ("Receive the goods", "Confirm received quantities. Receipt increases inventory at the selected branch."),
    ("Record supplier payment", "Capture payments made to the supplier."),
    ("Track the balance", "Any unpaid purchase amount remains visible as a supplier balance."),
)
add_table(doc, ["Step", "Branch action", "System result"], [[str(i + 1), name, result] for i, (name, result) in enumerate(purchase_steps)], [Inches(0.55), Inches(1.55), Inches(4.55)])
doc.add_paragraph(
    "Creating or ordering a purchase does not increase inventory. Stock changes only when an authorised employee receives the goods."
)

add_heading(doc, "Inventory Control", 1)
add_table(
    doc,
    ["Activity", "When the branch uses it", "Effect"],
    [
        ["Opening stock", "When the branch first begins tracking a product variant.", "Creates the starting inventory balance."],
        ["Stock adjustment", "When a documented correction, damage, loss, or other authorised change is required.", "Posts an increase or decrease with a reason and audit record."],
        ["Physical stock count", "During scheduled or investigative counts.", "Locks affected items during the count and reconciles differences at completion."],
        ["Minimum stock level", "When management defines the reorder threshold for an item.", "Drives low-stock and out-of-stock monitoring."],
        ["Stock movement history", "When employees need to explain a balance.", "Shows each receipt, sale, transfer, adjustment, and count-related movement."],
    ],
    [Inches(1.35), Inches(2.85), Inches(2.45)],
)

add_heading(doc, "Inter Branch Stock Transfers", 1)
transfer_intro = doc.add_paragraph(
    "Transfers move product variants from a source branch to a different active destination branch. The process protects both branch balances and identifies stock that is physically travelling between locations."
)
transfer_intro.paragraph_format.keep_with_next = True
transfer_steps = (
    ("Draft", "Source branch", "Create the transfer and add requested quantities. Inventory does not change."),
    ("Requested", "Source branch", "Submit the transfer for approval. Inventory still does not change."),
    ("Approved", "Source Branch Manager", "Approve the requested quantities. Inventory still does not change."),
    ("Dispatched", "Source branch", "Confirm dispatch. Stock leaves the source and becomes in transit."),
    ("Received", "Destination branch", "Confirm receipt. Stock enters the destination inventory."),
)
add_table(doc, ["Status", "Responsible location", "Inventory effect"], transfer_steps, [Inches(1.15), Inches(1.8), Inches(3.7)])
doc.add_paragraph(
    "An eligible transfer can be cancelled before dispatch. A dispatched transfer cannot be cancelled or edited. Dispatch and receipt apply to the complete transfer in the current workflow."
)

add_heading(doc, "Transfer Permissions by Role", 2)
add_table(
    doc,
    ["Role", "Read", "Create", "Approve", "Dispatch", "Receive", "Cancel"],
    [
        ["Super Admin", "All", "Yes", "Yes", "Yes", "Yes", "Yes"],
        ["Owner", "All", "Yes", "Yes", "Yes", "Yes", "Yes"],
        ["Branch Manager", "Related", "Source", "Source", "Source", "Destination", "Source"],
        ["Inventory Officer", "Related", "Source", "No", "Source", "Destination", "No"],
        ["Sales or Cashier", "Related", "No", "No", "No", "No", "No"],
    ],
    [Inches(1.2), Inches(0.75), Inches(0.8), Inches(0.8), Inches(0.85), Inches(0.9), Inches(0.8)],
)

add_heading(doc, "Shared and Branch Specific Information", 1)
add_table(
    doc,
    ["Shared across the company", "Recorded with branch identity"],
    [
        ["Product catalogue and variants", "Inventory quantities and minimum stock levels"],
        ["Categories, brands, units, and attributes", "Sales and sale payments"],
        ["Customer master records", "Purchases and supplier payments"],
        ["Supplier master records", "Receivables and supplier balances"],
        ["Public website catalogue", "Stock movements, adjustments, and counts"],
        ["System roles and permissions", "Transfers, branch dashboards, reports, and operational audit events"],
    ],
    [Inches(3.2), Inches(3.45)],
)
doc.add_paragraph(
    "For example, the same product can exist in the shared catalogue while Branch A holds 20 units and Branch B holds 4 units. Each branch sees its own quantity unless the employee has company-wide access."
)

add_heading(doc, "Dashboard and Reporting", 1)
doc.add_paragraph(
    "The dashboard adapts to the signed-in employee. Company roles receive company measures and branch comparisons. Branch employees receive information limited to their assignments and responsibilities."
)
add_table(
    doc,
    ["User", "Dashboard emphasis", "Reporting access"],
    [
        ["Branch Manager", "Branch sales, collections, purchasing, receivables, inventory health, transfers, and activity.", "Authorised branch reports and export."],
        ["Sales or Cashier", "Sales, customer collections, and receivable operations.", "Operational dashboard only under the standard role."],
        ["Inventory Officer", "Inventory health, urgent stock items, purchase receipt, and transfers.", "Operational dashboard only under the standard role."],
        ["Owner or Management", "Company totals, branch rankings, trends, inventory, receivables, purchasing, and transfers.", "Company and branch reports with export."],
        ["Super Admin", "The same company oversight plus administrative access.", "Company and branch reports with export."],
    ],
    [Inches(1.4), Inches(3.55), Inches(1.7)],
)

add_heading(doc, "Head Office Responsibilities", 1)
doc.add_paragraph("Owner or Management users monitor the company through:")
for item in (
    "Company and branch sales, collections, and purchasing results.",
    "Branch comparisons and product performance.",
    "Inventory levels, low-stock items, and stock in transit.",
    "Customer receivables and supplier balances.",
    "Employee access, roles, and audit history where authorised.",
    "Public website content and quotation requests.",
):
    add_bullet(doc, item)
doc.add_paragraph(
    "The Super Admin additionally creates and updates branches, provisions employee access, assigns roles, and manages branch assignments."
)

add_heading(doc, "Security and Audit Controls", 1)
doc.add_paragraph(
    "An authentication account alone does not grant application access. The employee also needs an active profile, a valid role, the application access permission, and an active branch assignment for branch-scoped work."
)
add_table(
    doc,
    ["Control", "Protection provided"],
    [
        ["Permission-aware interface", "Shows only the modules and actions available to the role."],
        ["Server authorisation", "Checks the employee's permission and branch before processing an operation."],
        ["Database row security", "Prevents direct access to operational records outside the employee's authorised scope."],
        ["Transaction controls", "Locks and validates critical inventory, sale, purchase, and transfer operations."],
        ["Audit history", "Records the responsible user, action, affected record, branch, and time for important events."],
        ["Inactive-account enforcement", "Stops an inactive employee from using application operations."],
    ],
    [Inches(1.75), Inches(4.9)],
)
doc.add_paragraph(
    "Typing another branch identifier into a URL or request does not grant access because server and database checks independently enforce the assignment."
)

add_heading(doc, "Recommended Daily Branch Routine", 1)
add_heading(doc, "Opening the Branch", 2)
for item in (
    "Confirm that staff can sign in with their individual accounts.",
    "Review dashboard alerts, transfers awaiting action, purchases awaiting receipt, and urgent stock items.",
    "Confirm the branch context before recording transactions.",
):
    add_bullet(doc, item)
add_heading(doc, "During the Day", 2)
for item in (
    "Complete sales when goods leave the branch and record payments against the correct customer.",
    "Receive supplier deliveries only after checking the physical quantities.",
    "Record stock corrections through authorised adjustments with clear reasons.",
    "Dispatch and receive transfers at the actual physical handover points.",
):
    add_bullet(doc, item)
add_heading(doc, "Closing the Branch", 2)
for item in (
    "Review completed sales, collections, unpaid balances, and exceptions.",
    "Confirm that received goods and transfers were entered correctly.",
    "Investigate unexpected stock balances through the movement history.",
    "Escalate unresolved differences to the Branch Manager and retain supporting records.",
):
    add_bullet(doc, item)

add_heading(doc, "Practical Example", 1)
doc.add_paragraph(
    "A customer buys five units of a product at Branch A and pays part of the amount. The cashier completes the sale. Branch A stock decreases by five units, the payment is recorded as a collection, and the unpaid amount becomes a receivable. Management can see the updated sale and receivable in authorised reports. Branch B stock remains unchanged because stock is maintained separately by branch."
)
doc.add_paragraph(
    "If Branch B later needs this product, an authorised employee at Branch A creates a transfer. Approval alone does not move stock. Dispatch removes the confirmed quantity from Branch A and reports it as in transit. Receipt at Branch B adds the quantity to Branch B inventory and completes the transfer."
)

add_heading(doc, "Key Operating Rules", 1)
for item in (
    "Use individual employee accounts and keep branch assignments current.",
    "Complete transactions in the branch where the physical activity occurs.",
    "A purchase changes inventory only when goods are received.",
    "A sale changes inventory when the sale is completed.",
    "A transfer removes source stock at dispatch and adds destination stock at receipt.",
    "Record payments against the correct customer, supplier, sale, purchase, and branch.",
    "Use adjustments only for documented corrections and complete physical counts before reconciling differences.",
    "Review dashboards and reports regularly rather than waiting for month end.",
):
    add_bullet(doc, item)

doc.core_properties.title = "SAT J Ent Branch Operations Guide"
doc.core_properties.subject = "How the SAT-J Ent system is used from each branch"
doc.core_properties.author = "SAT-J Ent"
doc.core_properties.keywords = "SAT-J Ent, branches, roles, sales, purchasing, inventory, transfers, security"
doc.save(OUTPUT)
print(OUTPUT)
