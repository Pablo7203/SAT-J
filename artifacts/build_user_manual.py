from pathlib import Path
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "SAT-J_Ent_User_Manual.docx"
NAVY, BLUE, PALE, GRAY, LINE = "17365D", "2F5597", "EDF3F8", "666666", "D9D9D9"

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); el = pr.find(qn("w:shd"))
    if el is None: el = OxmlElement("w:shd"); pr.append(el)
    el.set(qn("w:fill"), fill)

def margins(cell, n=110):
    pr = cell._tc.get_or_add_tcPr(); mar = pr.first_child_found_in("w:tcMar")
    if mar is None: mar = OxmlElement("w:tcMar"); pr.append(mar)
    for side in ("top", "start", "bottom", "end"):
        el = mar.find(qn(f"w:{side}"))
        if el is None: el = OxmlElement(f"w:{side}"); mar.append(el)
        el.set(qn("w:w"), str(n)); el.set(qn("w:type"), "dxa")

def borders(table):
    pr = table._tbl.tblPr; bs = pr.first_child_found_in("w:tblBorders")
    if bs is None: bs = OxmlElement("w:tblBorders"); pr.append(bs)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = bs.find(qn(f"w:{edge}"))
        if el is None: el = OxmlElement(f"w:{edge}"); bs.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:color"), LINE)

def no_split(row):
    pr = row._tr.get_or_add_trPr()
    if pr.find(qn("w:cantSplit")) is None: pr.append(OxmlElement("w:cantSplit"))

def repeat(row):
    pr = row._tr.get_or_add_trPr(); el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); pr.append(el)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False; borders(t)
    repeat(t.rows[0]); no_split(t.rows[0])
    for i, text in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,NAVY); margins(c); c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
        p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(text); r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
        if widths: c.width=widths[i]
    for ri, values in enumerate(rows):
        row=t.add_row(); no_split(row)
        for i, value in enumerate(values):
            c=row.cells[i]; margins(c); c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            if ri%2: shade(c,PALE)
            if widths: c.width=widths[i]
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.add_run(str(value))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def heading(text, level=1):
    p=doc.add_heading(text, level); p.paragraph_format.keep_with_next=True; return p

def para(text, lead=None):
    p=doc.add_paragraph()
    if lead: p.add_run(lead).bold=True
    p.add_run(text); return p

def bullet(text, level=0):
    return doc.add_paragraph(text, style="List Bullet" if level==0 else "List Bullet 2")

def steps(items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.24)
        p.paragraph_format.first_line_indent = Inches(-.24)
        p.add_run(f"{index}.  ").bold = True
        p.add_run(item)

def procedure(title, purpose, items, verify=None, caution=None):
    heading(title, 2); para(purpose)
    steps(items)
    if verify: para(verify, "Check the result  ")
    if caution: para(caution, "Important  ")

def page():
    if not getattr(page, "used", False):
        doc.add_page_break()
        page.used = True

def page_num(p):
    r=p.add_run(); begin=OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"),"begin")
    instr=OxmlElement("w:instrText"); instr.set(qn("xml:space"),"preserve"); instr.text=" PAGE "
    sep=OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"),"separate")
    end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end"); r._r.extend([begin,instr,sep,end])

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.82); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
normal=doc.styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(10.3); normal.font.color.rgb=RGBColor(35,35,35); normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.1
for name,size,before,after in (("Title",25,0,12),("Heading 1",17,14,7),("Heading 2",13,10,5),("Heading 3",11,8,4)):
    s=doc.styles[name]; s.font.name="Aptos Display"; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
    b=s._element.get_or_add_pPr().find(qn("w:pBdr"))
    if b is not None: s._element.get_or_add_pPr().remove(b)
for name in ("List Bullet","List Bullet 2","List Number"):
    s=doc.styles[name]; s.font.name="Aptos"; s.font.size=Pt(10.2); s.paragraph_format.space_after=Pt(3); s.paragraph_format.line_spacing=1.06
foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=foot.add_run("SAT-J Ent User Manual   |   "); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY); page_num(foot)

title=doc.add_paragraph(style="Title"); title.add_run("SAT J Ent User Manual")
sub=doc.add_paragraph(); rr=sub.add_run("Detailed operating guide for administrators management and branch employees"); rr.bold=True; rr.font.size=Pt(13); rr.font.color.rgb=RGBColor.from_string(GRAY)
sub.paragraph_format.space_after=Pt(14)
para("This manual explains how to use the SAT-J Ent system safely and consistently across all branches. It covers every main application area, the responsibilities of each role, the effect of each transaction, and the checks users should make before and after posting. Use it for onboarding, daily work, supervision, and issue investigation.", "Purpose  ")
para("The instructions describe the system as deployed to staging in September 2026. A production address and organization-specific support contact can be inserted before operational rollout.", "Document scope  ")
table(["Document item","Value"],[["System","SAT-J Ent"],["Audience","Super Admin, Owner, Branch Manager, Sales or Cashier, Inventory Officer"],["Coverage","Administration, catalogue, inventory, purchasing, sales, transfers, reporting, website, quotations"],["Version","1.0 - September 2026"]],[Inches(1.55),Inches(5.1)])

heading("How to Use This Manual",1)
bullet("New employees should read Getting Started, Roles and Access, and the section for their assigned role before using live data.")
bullet("Branch Managers should also read the daily controls, exception handling, and reporting sections.")
bullet("Super Admin and Owner users should read the administration, security, audit, and public website sections.")
bullet("Words such as complete, receive, dispatch, reconcile, cancel, reverse, and archive describe controlled actions with lasting operational or audit effects.")

heading("Contents",1)
for x in ("1 Getting Started and Navigation","2 Roles Branch Scope and Security","3 Dashboard and Profile","4 Branch Employee and Access Administration","5 Catalogue Setup Products and Pricing","6 Inventory Operations","7 Suppliers Purchases and Supplier Payments","8 Customers Sales Receivables and Receipts","9 Inter Branch Stock Transfers","10 Quotations Public Website and Catalogue","11 Reports Exports and Reconciliation","12 Daily Branch Procedures and Controls","13 Troubleshooting and Support","14 Quick Reference"):
    bullet(x)

page(); heading("1 Getting Started and Navigation",1)
heading("System Operating Model",2)
para("All branches use one web application and one central database. Shared records include products, variants, customers, suppliers, categories, brands, units, attributes, and system roles. Operational records retain their branch identity, including inventory, sales, purchases, payments, transfers, dashboards, and reports.")
table(["Concept","Meaning for the user"],[["Role","Controls the modules and actions available to the employee."],["Branch assignment","Limits branch-scoped employees to active assigned branches."],["Status","Controls which next actions are allowed for a record."],["Audit history","Records important administrative and operational actions."],["Immutable evidence","Completed stock movements, receipts, and payments are retained rather than silently edited."]],[Inches(1.4),Inches(5.25)])

procedure("Sign In","Use the individual account provided by an administrator.",["Open the organization-provided SAT-J Ent login address.","Enter the employee email address and password.","Select Sign in.","After successful authentication, confirm that the dashboard and sidebar match the employee's role."],"The employee name, role, and authorized modules should be visible. An inactive or unassigned branch account should not gain application access.","Do not share accounts or save passwords on public devices. The audit trail identifies the signed-in account.")
procedure("Sign Out","End the authenticated session before leaving a shared device.",["Open the account or sidebar controls.","Select Sign out.","Confirm that the login page appears before leaving the device."],"Protected application pages should require a new sign-in.")
procedure("Reset a Forgotten Password","Request a recovery link when email delivery is configured.",["Open Forgot password from the login page.","Enter the employee email address and submit.","Open the recovery email and follow the link before it expires.","Enter and confirm a new strong password on Reset password.","Return to login and sign in with the new password."],"The request screen intentionally gives a neutral response even when an address is unknown.","If no email arrives, contact the Super Admin. Staging or production SMTP must be configured for delivery.")

heading("Desktop and Mobile Navigation",2)
para("On a laptop or desktop, use the left sidebar. The active page is highlighted. On a phone, open the hamburger menu, select a page, and the menu closes automatically. Page actions may stack vertically on smaller screens.")
bullet("Use search and filters before scrolling through long lists.")
bullet("Use the browser Back button only when a page does not provide a clearer application action.")
bullet("Avoid double-clicking Save, Complete, Dispatch, Receive, or Payment buttons. Wait for the result message.")
bullet("If a filtered dashboard or report seems unchanged, select Apply filters and confirm the date and branch.")

page(); heading("2 Roles Branch Scope and Security",1)
table(["Role","Primary use","Scope"],[["Super Admin","Branches, employees, access, full operations, audit, configuration","Company"],["Owner","Company oversight, catalogue, all operations, reporting, website","Company"],["Branch Manager","Full branch operations and branch reporting","Assigned branches"],["Sales or Cashier","Customers, sales, payments, receivables, quotations","Assigned branches"],["Inventory Officer","Stock, receipts, counts, adjustments, transfers","Assigned branches"]],[Inches(1.35),Inches(3.95),Inches(1.35)])
heading("Access Rules",2)
bullet("An Auth account alone does not grant access. The user also needs an active employee profile, role, app access permission, and active branch assignment when branch-scoped.")
bullet("Company roles can work across branches. Branch roles can work only inside active assignments.")
bullet("The interface, server actions, and database row security all enforce access independently.")
bullet("Typing another branch or record identifier into a URL does not grant access.")
bullet("Deactivation prevents protected work even if an old session token has not yet expired.")
heading("Responsibility Boundaries",2)
table(["Activity","Typical responsible role","Required control"],[["Manage employees and branches","Super Admin","Verify identity, role, branch assignment, and active state."],["Approve sensitive branch operations","Branch Manager or company role","Review status, branch, quantities, values, and supporting evidence."],["Record customer transactions","Sales or Branch Manager","Use correct customer, branch, price, payment, and receipt."],["Record stock changes","Inventory or Branch Manager","Use receipts, sales, transfers, adjustments, or counts; never direct edits."],["Review company results","Owner or Super Admin","Reconcile status, date range, timezone, branch, and underlying transactions."]],[Inches(1.65),Inches(2.2),Inches(2.85)])

page(); heading("3 Dashboard and Profile",1)
heading("Dashboard",2)
para("The dashboard provides an operational snapshot for a selected date range and branch. Company users see executive totals and branch comparisons. Branch roles see measures appropriate to their responsibilities and assignments.")
steps(["Confirm the From and To dates.","Choose All branches or one authorized branch where available.","Select Apply filters.","Review the summary cards before the detailed sections.","Investigate unexpected figures through the relevant module or report rather than correcting totals directly."])
table(["Dashboard area","Interpretation"],[["Sales revenue","Value of completed sales in the selected period."],["Completed sales and units sold","Completed operational volume, not draft activity."],["Customer collections","Posted customer payments; this is separate from sales revenue."],["Purchase value","Non-cancelled purchasing value, not cash paid."],["Inventory health","Low-stock, out-of-stock, and in-transit indicators."],["Receivable aging","Outstanding named-customer completed sale balances."],["Branch performance","Comparison of completed-sales revenue by branch."],["Transfer operations","Requested, approved, dispatched, and in-transit activity."]],[Inches(1.75),Inches(4.9)])
para("Revenue and collections answer different questions. A credit sale increases revenue at completion but does not increase collections until payment is posted.","Interpretation rule  ")
heading("My Profile",2)
para("My Profile shows the signed-in employee identity and current access configuration. Review it when expected branches or modules are missing. Employees cannot grant themselves permissions from this page.")

page(); heading("4 Branch Employee and Access Administration",1)
procedure("Create a Branch","Super Admin users create the operational location before assigning employees or posting stock.",["Open Branches.","Enter the unique branch code, name, address, phone, and other available details.","Save the branch.","Review the saved branch card.","Activate the branch when it is ready for use.","Enable public visibility only after verifying the contact information is suitable for the website."],"The branch should appear in administration and relevant operational selectors.","Deactivating a branch affects future access and operations. Preserve historical records.")
procedure("Invite an Employee","Provision one account per employee.",["Open Employees.","Select or open Invite employee.","Enter the employee email and full name.","Choose the stable system role.","For Branch Manager, Sales, or Inventory, select at least one active branch.","Confirm the invitation.","Verify the employee profile, role, assignments, and active state."],"The employee should appear in Employees. Email delivery depends on configured SMTP.","Do not use placeholder addresses for real production employees. Never enable public employee signup.")
procedure("Change Employee Access","Use this for transfers, promotion, role change, temporary suspension, or departure.",["Open Employees and locate the person.","Review the current role and assignments.","Select the new role and authorized branches.","Set the active state deliberately.","Save and confirm the updated access summary.","Review the access audit history for the recorded change."],"The employee should immediately be limited to the new configuration.","Self-management and removal of the final active Super Admin are blocked.")
heading("Roles and Access Page",2)
para("Roles and Access is a read-only role explorer. Filter by role, permission, or module to understand what a role can do. Assign roles from Employees; do not attempt to edit system roles.")
heading("Access Audit History",2)
para("The audit page lists recent access-control and branch-administration events. Use it to answer who changed access, what entity was affected, when the change occurred, and which branch was involved. Application users cannot edit or delete audit records.")

page(); heading("5 Catalogue Setup Products and Pricing",1)
heading("Recommended Setup Order",2)
steps(["Create categories.","Create brands where applicable.","Create units of measure.","Create product attributes and allowed values.","Map attributes to categories and mark required attributes.","Create products with their default variants.","Add additional variants and attribute values.","Set retail and wholesale prices.","Add images and public visibility settings.","Activate the product only after validation succeeds."])
heading("Categories Brands and Units",2)
para("These are company-wide reference records. Inactive records remain visible to historical products. Use clear names, stable codes, and consistent units. Avoid creating near-duplicates with spelling differences.")
heading("Attributes",2)
table(["Type","Use","Example"],[["Text","Free-form characteristic","Model code"],["Number","Numeric characteristic","Thickness"],["Boolean","Yes or no characteristic","Slip resistant"],["Select","Controlled allowed values","Finish: Matte or Gloss"]],[Inches(1.1),Inches(3.65),Inches(1.9)])
procedure("Create a Product","Create the descriptive product and required default sellable variant together.",["Open Products and select New product.","Enter the product name, category, optional brand, unit of measure, and description.","Enter the default variant name, unique SKU, and optional barcode.","Save the product in Draft status.","Open Edit and complete required category attributes.","Add further variants where the same product has distinct sellable options.","Add current prices and approved images.","Select Activate only when the record is operationally ready."],"The product detail should show the expected variants, attributes, prices, and Draft or Active status.","SKU and barcode values must be unique. Activation requires an active category and unit, an active SKU variant, and all required attributes.")
procedure("Change a Price","Create effective-dated price history rather than overwriting old values.",["Open Catalogue setup and Price history or open the relevant product.","Choose the variant, Retail or Wholesale type, and company or branch scope.","Enter the new GHS amount and effective date.","Save the price change.","Confirm that the old applicable record has an end date and the new current record does not."],"Only one price should apply for the same variant, scope, type, and time.","Overlapping price periods are rejected. Branch-specific prices take operational scope into account.")
heading("Product Lifecycle",2)
table(["Status","Meaning","Normal next action"],[["Draft","Being prepared; not operationally public","Complete data and activate."],["Active","Available for authorized operations","Maintain variants, pricing, images, and visibility."],["Archived","Retained for history but not normal new use","Restore only after review."]],[Inches(1.05),Inches(3.55),Inches(2.05)])

page(); heading("6 Inventory Operations",1)
heading("Inventory Concepts",2)
table(["Field","Meaning"],[["Quantity on hand","Current branch quantity derived from controlled stock movements."],["Minimum level","Branch reorder or attention threshold."],["Available source stock","Advisory quantity for an operation; final posting rechecks stock."],["In transit","Dispatched transfer stock not yet received at the destination."],["Movement history","Immutable evidence showing type, signed quantity, before balance, after balance, and source document."]],[Inches(1.65),Inches(5.0)])
procedure("Post Opening Stock","Initialize a variant at a branch once.",["Open Inventory and select Opening stock.","Choose the active branch.","Add product variants by SKU and enter verified starting quantities.","Optionally enter minimum levels.","Enter a clear confirmation note.","Review all lines against the approved opening count.","Post opening stock."],"Balances should appear in Inventory with OPENING_STOCK movements.","Do not repeat opening stock to correct a mistake. Use an approved Data Correction adjustment.")
procedure("Update a Minimum Level","Maintain the branch-specific threshold used by stock alerts.",["Open Inventory and select the relevant branch and variant.","Open the inventory detail.","Enter the approved minimum level.","Select Update minimum.","Confirm the new value and resulting low-stock classification."],"Dashboard and inventory alerts should reflect the threshold.")
procedure("Create and Complete a Stock Adjustment","Use adjustments for documented damage, loss, found stock, or correction.",["Open Inventory, Stock adjustments, and New adjustment.","Choose the branch and adjustment reason.","Enter a meaningful note.","Add variants and signed or direction-appropriate quantities.","Save the draft.","Review available stock and every line.","Select Complete adjustment."],"The document should become Completed and matching movements should update balances.","Only completion changes stock. If any line is invalid, the entire posting rolls back. Correct completed work with another adjustment.")
procedure("Perform a Physical Stock Count","Reconcile physical quantities against a controlled snapshot.",["Open Inventory, Stock counts, and New count.","Choose the branch and variants to count.","Create the draft count.","Select Start count to capture system quantities.","Count the physical stock and enter each quantity.","Save the count entries.","Review signed variances and investigate unusual differences.","Select Complete and reconcile count."],"The count should become Completed and variance movements should reconcile each included item.","Affected variants are protected from conflicting manual mutations during an active count. Restart or investigate when a concurrency warning appears.")
heading("Investigate a Stock Balance",2)
steps(["Open the inventory detail for the branch and variant.","Compare quantity on hand and minimum level.","Review Recent movements or open Stock movements.","Trace receipts, completed sales, transfers, adjustments, and count reconciliation.","Compare source documents and physical evidence.","Use a new approved adjustment only when a documented correction is justified."])
para("Never directly edit quantity on hand, insert or change ledger rows, delete movements, or use a service-role client for ordinary operations.","Prohibited action  ")

page(); heading("7 Suppliers Purchases and Supplier Payments",1)
procedure("Add a Supplier","Create a reusable company supplier record.",["Open Suppliers.","Search for similar names and contact details.","Select Add supplier.","Enter name, company, contact person, phone, email, address, and notes as available.","Save and review the supplier detail."],"The supplier should be available when creating purchases.","Similar names are allowed and are not silently merged. Do not delete historical suppliers; archive only when appropriate.")
procedure("Create and Order a Purchase","Record what an authorized branch has ordered.",["Open Purchases and select New purchase.","Select the supplier and authorized destination branch.","Enter purchase date, expected date, supplier invoice reference, and notes.","Add variants, ordered quantities, and unit costs.","Enter approved discount and other costs where applicable.","Save the purchase draft.","Review the database-calculated total.","Select Mark ordered when the supplier order has actually been placed."],"The status should be Ordered and inventory should remain unchanged.","Creating or ordering a purchase does not add stock.")
procedure("Receive Purchased Goods","Post only goods physically received at the selected branch.",["Open the ordered or partially received purchase.","Select Receive goods.","Compare ordered, previously received, and remaining quantities.","Inspect the physical delivery.","Enter only quantities present in this delivery.","Enter the delivery reference and receipt notes.","Confirm receipt."],"A goods receipt should be listed, purchase status should update, and branch inventory should increase.","The receipt is atomic and immutable. Repeat for later partial deliveries. Do not edit inventory to correct a receipt.")
procedure("Record a Supplier Payment","Record cash movement separately from ordering and receipt.",["Open the purchase and select Record payment.","Confirm supplier, branch, purchase, and outstanding balance.","Enter the payment amount, date, method, reference, and notes.","Submit once.","Review the supplier payment history and remaining balance."],"Payment status should become Partially Paid or Paid as appropriate.","The database blocks duplicate operation submissions and overpayment.")
heading("Purchase Status Guide",2)
table(["Status","Meaning"],[["Draft","Editable planned purchase; no stock effect."],["Ordered","Placed with supplier; no stock effect."],["Partially Received","Some goods posted; remaining lines can be received later."],["Received","All ordered quantities received."],["Cancelled","Eligible purchase stopped; retained for history."]],[Inches(1.55),Inches(5.1)])

page(); heading("8 Customers Sales Receivables and Receipts",1)
procedure("Add a Customer","Use a named customer whenever a balance may remain.",["Open Customers and search by name, phone, or email.","Select Add customer.","Choose the customer type.","Enter name, company, contact details, address, and notes.","Save and review the customer detail."],"The customer should be selectable for a sale.","The starred Walk-In Customer is a protected system record and cannot carry an unpaid balance.")
procedure("Create a Sale","Prepare a branch sale before completion changes inventory.",["Open Sales and select New sale.","Select the authorized branch and customer.","For credit, choose a named customer and enter a due date where appropriate.","Add product variants and quantities.","Review resolved prices.","Enter an authorized discount or price override only with the required reason.","Save the sale draft.","Review items, subtotal, discount, total, customer, and branch."],"The sale should remain Draft and stock should remain unchanged.")
procedure("Complete a Fully Paid Sale","Complete the sale when goods leave the branch and full payment is received.",["Open the reviewed draft sale.","Select Complete sale.","Enter the full displayed amount.","Choose payment method and enter a reference or note where required.","Confirm once.","Open Print receipt and provide the customer copy."],"The sale should be Completed and Paid; branch stock should decrease and a posted payment should appear.","Stock is checked again at completion. If any item is short, nothing is completed.")
procedure("Complete a Credit or Partial Payment Sale","Record revenue and inventory while preserving the customer balance.",["Use a named customer, never Walk-In.","Review the draft and due date.","Select Complete sale.","Enter zero for full credit or less than the total for partial payment.","Enter payment method details only when an initial payment is recorded.","Confirm and review the resulting balance."],"The sale should show Unpaid or Partially Paid and appear in Receivables.")
procedure("Record a Later Customer Payment","Apply a collection to an outstanding completed sale.",["Open the outstanding sale or customer detail.","Select Record payment.","Confirm customer, sale, branch, and balance.","Enter an amount no greater than the balance.","Choose method, date, reference, and notes.","Submit and review the payment history."],"The balance and payment status should update. Collections reporting should include the posted payment.")
procedure("Reverse an Incorrect Customer Payment","Preserve the original record and post a controlled reversal.",["Open the sale or customer payment history.","Locate the incorrect posted payment.","Enter a specific reversal reason.","Select Reverse and confirm.","Record the correct replacement payment separately if needed."],"The original and reversal evidence should remain visible, and the sale balance should be recalculated.","Only authorized management can reverse payments. Never edit or delete a payment.")
procedure("Cancel an Erroneous Completed Sale","Reverse an eligible sale while preserving evidence.",["Open the completed sale.","Reverse every posted payment first.","Confirm that cancellation is operationally correct and not a return or exchange.","Enter the cancellation reason.","Select Cancel and confirm."],"The sale should become Cancelled and inventory should be restored through SALE_REVERSAL movements.","Do not use cancellation for genuine returns or exchanges; those require an approved business process.")
heading("Receivables",2)
para("Receivables lists positive balances for completed named-customer sales within the authorized branch scope. Filter for overdue balances or records without a due date. Follow up using the customer and sale details, then record payments against the correct sale.")

page(); heading("9 Inter Branch Stock Transfers",1)
table(["Status","Responsible location","Stock effect"],[["Draft","Source","None"],["Requested","Source","None"],["Approved","Source manager","None; approval does not reserve stock"],["Dispatched","Source","TRANSFER_OUT; quantity becomes in transit"],["Received","Destination","TRANSFER_IN; destination stock increases"],["Cancelled","Source before dispatch","None"]],[Inches(1.15),Inches(1.8),Inches(3.7)])
procedure("Create and Submit a Transfer","Request full quantities from an authorized source to a different active destination.",["Open Stock transfers and select New transfer.","Choose the source branch and different destination branch.","Review Available source stock.","Add variants and requested quantities.","Enter transfer notes.","Save draft for later editing or save and submit."],"Draft or Requested status should appear. Inventory should not change.")
procedure("Approve a Transfer","Authorize a requested transfer at the source branch.",["Open the Requested transfer.","Confirm source, destination, requested variants, quantities, and business need.","Select Approve requested quantities.","Confirm the status changes to Approved."],"Stock remains unchanged and unreserved until dispatch.")
procedure("Dispatch a Transfer","Record the physical handover from the source branch.",["Open the Approved transfer at the source.","Recount the complete transfer and verify destination documents.","Select Confirm full dispatch.","Wait for the successful result before handing over the final documentation."],"Status should become Dispatched, source stock should decrease, and quantities should appear in transit.","Current stock is rechecked atomically. Partial dispatch is not supported.")
procedure("Receive a Transfer","Post the physical delivery at the destination branch.",["Open the Dispatched transfer at the destination.","Compare every line and quantity with the physical delivery.","If complete and undamaged, select Confirm full receipt.","Confirm the status becomes Received."],"Destination inventory should increase and in-transit quantity should clear.","If goods are short or damaged, do not receive and do not adjust stock manually. Keep the transfer in transit and escalate.")
heading("Transfer Permission Summary",2)
table(["Role","Create","Approve","Dispatch","Receive","Cancel"],[["Super Admin or Owner","Yes","Yes","Yes","Yes","Yes"],["Branch Manager","Source","Source","Source","Destination","Source"],["Inventory Officer","Source","No","Source","Destination","No"],["Sales or Cashier","No","No","No","No","No"]],[Inches(1.55),Inches(1.0),Inches(1.0),Inches(1.0),Inches(1.0),Inches(1.0)])

page(); heading("10 Quotations Public Website and Catalogue",1)
heading("Public Catalogue",2)
para("Visitors can browse approved active public products, categories, brands, variants, attributes, and selected prices without employee access. Products that are Draft, Archived, or private must not appear in listings, direct lookup, sitemap, or anonymous image access.")
procedure("Publish or Hide a Product","Control whether a catalogue product is visible to the public.",["Complete the product, variants, required attributes, prices, and images.","Activate the product.","Open Edit product.","Enable Publicly visible.","Enable Show retail price online only for an approved public retail price.","Optionally enable Feature on public website.","Save and verify the public product page.","To hide it, remove public visibility or return it to Draft or Archived status."],"Anonymous catalogue and direct product behavior should match the chosen visibility.","Descriptions and images must not disclose supplier terms, cost, margin, or internal notes.")
procedure("Manage Website Contact and Hero Content","Maintain accurate public company information.",["Open Public website in administration.","Review hero heading and supporting copy.","Enter verified phone, email, and WhatsApp values.","Save settings.","Open the public site and test visible contact actions."],"The public site should show the approved content.","Leave unavailable fields empty rather than publishing placeholders. WhatsApp numbers should include the country code.")
procedure("Process a Quotation Request","Move a public enquiry through its follow-up lifecycle.",["Open Quotations.","Filter or select a request within authorized scope.","Review prospect contact, company, product or project request, quantity, branch, source, and message.","Contact the prospect outside the system using approved channels.","Update status to Contacted.","Use In Progress while preparing or following up.","Set Closed when resolved or Cancelled when discontinued."],"The request detail should show the new status and relevant timestamps.","A quotation request does not reserve stock, create a customer, or create a sale automatically.")

page(); heading("11 Reports Exports and Reconciliation",1)
heading("Using Reports",2)
steps(["Open Reports.","Choose the report type.","Set From and To dates; ranges are limited to two years.","Choose an authorized branch or company scope where available.","Apply filters.","Review totals and detailed rows.","Use Export only when permitted and needed; detailed exports are limited to 1,000 rows.","Store exported files according to company data-handling policy."])
table(["Measure","Reconcile to"],[["Sales revenue","Completed sale totals using completed time in Africa/Accra."],["Collections","Posted customer payments; do not equate to revenue."],["Purchase value","Non-cancelled purchase totals."],["Supplier payments","Posted supplier payment records; separate from purchase value."],["Receivable aging","Positive balances on completed named-customer sales."],["Inventory","Branch balances and immutable movement history."],["Transfers in transit","Dispatched transfers not yet received."]],[Inches(1.75),Inches(4.9)])
heading("When a Figure Looks Wrong",2)
steps(["Confirm the selected date range.","Confirm branch or All branches.","Confirm the record status required by the measure.","Confirm whether the measure uses transaction date, completion date, or payment date.","Remember that reporting uses the Africa/Accra business timezone.","Open the underlying sales, payments, purchases, movements, or transfers.","Document the discrepancy and escalate if source records cannot explain it."])
para("Never use a service-role session to diagnose what a normal employee can see; it bypasses the employee's caller context.","Security rule  ")

page(); heading("12 Daily Branch Procedures and Controls",1)
heading("Opening the Branch",2)
for x in ("Confirm employees can sign in using individual accounts.","Review dashboard alerts and branch/date filters.","Review transfers awaiting approval, dispatch, or receipt.","Review purchases awaiting receipt and urgent stock items.","Confirm device, printer, and network readiness where required.","Confirm the correct branch before the first transaction."): bullet(x)
heading("During the Day",2)
for x in ("Complete sales when goods physically leave the branch.","Record payments against the correct customer, sale, purchase, supplier, and branch.","Receive supplier deliveries only after physical verification.","Dispatch and receive transfers at the actual handover points.","Use approved stock adjustments with clear reasons.","Protect customer and employee information from unauthorized viewing."): bullet(x)
heading("Closing the Branch",2)
for x in ("Review completed sales, collections, credit balances, and exceptions.","Confirm supplier receipts and payments entered during the day.","Review transfers still in transit or awaiting action.","Investigate unexpected balances through movement history.","Escalate unresolved differences and retain supporting evidence.","Sign out of shared devices."): bullet(x)
heading("Period End Controls",2)
table(["Frequency","Control"],[["Daily","Review sales, payments, exceptions, pending receipts, and transfers."],["Weekly","Review receivables, supplier balances, low stock, and overdue operational queues."],["Monthly","Reconcile dashboard and reports to transaction detail; perform approved counts."],["Quarterly","Review employee access, branch assignments, inactive accounts, catalogue quality, and website content."],["As needed","Rotate compromised credentials, investigate audit events, and document corrections."]],[Inches(1.2),Inches(5.45)])

doc.add_page_break(); heading("13 Troubleshooting and Support",1)
table(["Problem","Checks and response"],[["Cannot sign in","Confirm URL, email spelling, password, network, and account status. Use password recovery if SMTP is configured."],["Access denied","Check active profile, role, permission, and active branch assignment. Do not bypass security."],["Expected branch missing","Ask Super Admin to review employee branch assignments and branch active state."],["Save or completion rejected","Read the message; check required fields, status, stock, balance, branch, duplicate submission, and permission."],["Dashboard appears wrong","Reapply dates and branch, then compare underlying status and transaction dates."],["Stock balance unexpected","Trace immutable movements and source documents. Use an approved correction only after investigation."],["Purchase cannot be cancelled","A receipt may exist or the status may be ineligible. Preserve the record and escalate."],["Sale cannot be cancelled","Reverse posted payments first and confirm the case is not a return or exchange."],["Transfer cannot be received","Only the destination can receive; physical delivery must match and status must be Dispatched."],["Recovery or invitation email missing","Verify SMTP and Auth redirect configuration; do not claim delivery without evidence."],["Public product missing","Check Active status, public visibility, variants, images, and required attributes."],["Export denied or incomplete","Confirm reporting permission, branch scope, filters, and the 1,000-row detail limit."]],[Inches(2.05),Inches(4.6)])
heading("What to Include in a Support Request",2)
for x in ("Employee email and role, but never the password.","Page name and action attempted.","Branch and record number, such as sale, purchase, transfer, or adjustment number.","Approximate date and time with timezone.","Exact error message and a screenshot without unnecessary personal information.","Whether the issue is repeatable and whether another authorized user sees it.","Any physical document or operational evidence relevant to the transaction."): bullet(x)
heading("Security Incidents",2)
para("Immediately report suspected credential exposure, unexpected access, unexplained audit activity, or data visible outside authorized scope. Stop using the affected account, preserve evidence, deactivate or rotate credentials through an authorized administrator, and investigate before resuming work.")

page(); heading("14 Quick Reference",1)
heading("Transaction Effects",2)
table(["Action","Inventory effect","Financial or status effect"],[["Save purchase draft","None","Creates Draft purchase."],["Mark purchase ordered","None","Status becomes Ordered."],["Receive purchase","Increase at branch","Creates immutable goods receipt."],["Record supplier payment","None","Reduces purchase or supplier balance."],["Save sale draft","None","Creates Draft sale."],["Complete sale","Decrease at branch","Creates revenue, receipt, and customer balance."],["Record customer payment","None","Increases collections and reduces receivable."],["Dispatch transfer","Decrease at source","Creates in-transit stock."],["Receive transfer","Increase at destination","Completes transfer and clears in transit."],["Complete adjustment","Increase or decrease","Creates reasoned stock movements."],["Complete stock count","Variance increase or decrease","Reconciles snapshot to physical count."]],[Inches(1.85),Inches(1.65),Inches(3.15)])
heading("Status Rules to Remember",2)
for x in ("Draft records normally do not change inventory.","Purchase ordering does not change inventory; receipt does.","Sale completion changes inventory; recording a draft does not.","Transfer approval does not reserve or move stock; dispatch and receipt do.","Walk-In sales must be fully paid.","Completed evidence is corrected through reversal, cancellation, adjustment, or a new transaction, not silent editing.","Branch scope follows the employee's active assignment, not the branch name typed into a URL."): bullet(x)
heading("Administrator Go Live Checklist",2)
para("Before production use, verify hosting, Auth redirects, SMTP, employee access, approved business data, hosted security tests, backups, recovery readiness, and support ownership.")

doc.core_properties.title="SAT J Ent User Manual"
doc.core_properties.subject="Detailed operating guide for the SAT-J Ent business system"
doc.core_properties.author="SAT-J Ent"
doc.core_properties.keywords="SAT-J Ent, user manual, ERP, branches, inventory, sales, purchasing, transfers, reports"
doc.save(OUT)
print(OUT)
